from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from ar_mis.gate import GateStatus
from ar_mis.models import (
    CustomerMasterRecord,
    PTPEntry,
    PTPStatus,
    PTPStatusLogRow,
    WeeklySnapshotRow,
)
from ar_mis.reporting import generate_report
from ar_mis.storage import Store


def _row(party, branch, week, closing, reconciled=True):
    return WeeklySnapshotRow(
        party_id=party,
        branch_id=branch,
        week_ending=week,
        opening=Decimal("0.00"),
        sales=Decimal("0.00"),
        credit_notes=Decimal("0.00"),
        debit_notes=Decimal("0.00"),
        receipts=Decimal("0.00"),
        journals=Decimal("0.00"),
        closing_computed=Decimal(closing),
        closing_extracted=Decimal(closing) if reconciled else Decimal(closing) + Decimal("1.00"),
        reconciled=reconciled,
        difference=Decimal("0.00") if reconciled else Decimal("-1.00"),
    )


def _build_store(tmp_path):
    store = Store(str(tmp_path / "report.db"))
    store.upsert_customer_master(CustomerMasterRecord("Acme", "Acme Corp", "KOL", Decimal("0.00")))
    store.upsert_customer_master(CustomerMasterRecord("Acme", "Acme Corp", "DEL", Decimal("0.00")))
    store.upsert_customer_master(CustomerMasterRecord("Ghost", "Ghost Ltd", "KOL", Decimal("0.00")))

    week1 = date(2026, 1, 5)
    week2 = date(2026, 1, 12)
    store.append_weekly_snapshot(_row("Acme", "KOL", week1, "100000.00"))
    store.append_weekly_snapshot(_row("Acme", "DEL", week1, "50000.00"))
    store.append_weekly_snapshot(_row("Ghost", "KOL", week1, "20000.00"))
    # Week 2: Acme/KOL present, Acme/DEL present, but Ghost/KOL missing -
    # this is the gap the Movement Trend sheet must surface as MISSING.
    store.append_weekly_snapshot(_row("Acme", "KOL", week2, "110000.00"))
    store.append_weekly_snapshot(_row("Acme", "DEL", week2, "55000.00", reconciled=False))

    entry = PTPEntry("ptp-1", "Ghost", "KOL", week1, Decimal("20000.00"), date(2026, 1, 15), "AR Mgr")
    store.create_ptp_entry(entry)
    store.log_ptp_status(PTPStatusLogRow("ptp-1", week1, PTPStatus.ACTIVE, "AR Mgr"))

    return store, week1, week2


def test_generate_report_clean_gate(tmp_path):
    store, week1, _ = _build_store(tmp_path)
    output = tmp_path / "report.xlsx"
    generate_report(
        store, week1, GateStatus(clean=True, reasons=[]), final_failed_branches=[], drift_findings=[],
        output_path=str(output),
    )
    wb = load_workbook(output)
    assert set(wb.sheetnames) == {"Summary", "Party Detail", "Exceptions", "Movement Trend", "PTP Register"}

    summary = wb["Summary"]
    status_row = [r for r in summary.iter_rows(values_only=True) if r and r[0] == "Status"][0]
    assert status_row[1] == "VALIDATED"


def test_generate_report_unclean_gate_shows_reasons(tmp_path):
    store, week1, _ = _build_store(tmp_path)
    output = tmp_path / "report_unclean.xlsx"
    status = GateStatus(clean=False, reasons=["Branch 'Delhi' failed extraction after retry: timeout"])
    generate_report(
        store, week1, status, final_failed_branches=[], drift_findings=[], output_path=str(output)
    )
    wb = load_workbook(output)
    summary = wb["Summary"]
    values = [cell for row in summary.iter_rows(values_only=True) for cell in row if cell]
    assert any("UNVALIDATED" in str(v) for v in values)
    assert any("timeout" in str(v) for v in values)


def test_party_detail_includes_cross_branch_total(tmp_path):
    store, week1, _ = _build_store(tmp_path)
    output = tmp_path / "report.xlsx"
    generate_report(
        store, week1, GateStatus(clean=True, reasons=[]), final_failed_branches=[], drift_findings=[],
        output_path=str(output),
    )
    wb = load_workbook(output)
    detail = wb["Party Detail"]
    rows = list(detail.iter_rows(values_only=True))
    total_rows = [r for r in rows if r[1] == "TOTAL (all branches)"]
    assert len(total_rows) == 1
    # Acme: 100000 (KOL) + 50000 (DEL) = 150000
    assert total_rows[0][8] == "150000.00"


def test_exceptions_sheet_lists_reconciliation_failures_separately(tmp_path):
    store, week1, week2 = _build_store(tmp_path)
    output = tmp_path / "report.xlsx"
    generate_report(
        store, week2, GateStatus(clean=False, reasons=["x"]), final_failed_branches=[], drift_findings=[],
        output_path=str(output),
    )
    wb = load_workbook(output)
    exceptions = wb["Exceptions"]
    rows = list(exceptions.iter_rows(values_only=True))
    assert any(r[0] == "Party reconciliation FAIL" and "DEL" in r[1] for r in rows[1:])


def test_movement_trend_flags_missing_week_gap(tmp_path):
    store, week1, week2 = _build_store(tmp_path)
    output = tmp_path / "report.xlsx"
    generate_report(
        store, week2, GateStatus(clean=True, reasons=[]), final_failed_branches=[], drift_findings=[],
        output_path=str(output),
    )
    wb = load_workbook(output)
    trend = wb["Movement Trend"]
    rows = list(trend.iter_rows(values_only=True))
    ghost_week2_rows = [r for r in rows if r[0] == "Ghost" and r[2] == week2.isoformat()]
    assert len(ghost_week2_rows) == 1
    assert ghost_week2_rows[0][4] == "MISSING"


def test_ptp_register_shows_current_status(tmp_path):
    store, week1, _ = _build_store(tmp_path)
    output = tmp_path / "report.xlsx"
    generate_report(
        store, week1, GateStatus(clean=True, reasons=[]), final_failed_branches=[], drift_findings=[],
        output_path=str(output),
    )
    wb = load_workbook(output)
    ptp = wb["PTP Register"]
    rows = list(ptp.iter_rows(values_only=True))
    ptp_row = [r for r in rows if r[0] == "ptp-1"][0]
    assert ptp_row[6] == "Active"
