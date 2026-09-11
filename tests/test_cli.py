"""End-to-end smoke test for the actual script an operator runs each
Monday (ar_mis.cli.run). Everything downstream of TallyClient is real
(orchestration, pipeline, reconciliation, gate, reporting, SQLite
storage) - only the Tally transport itself is faked, since no live Tally
instance is reachable from this build environment (see README).
"""
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from ar_mis import cli
from ar_mis.config import BranchConfig
from ar_mis.models import CustomerMasterRecord, LedgerEntry, Voucher, VoucherType
from ar_mis.storage import Store

KOL = BranchConfig("KOL", "Kolkata", "Kolkata HQ")


class FakeTallyClient:
    """Stands in for ar_mis.tally_client.TallyClient. Always reports one
    Sales voucher for 'Acme' (Rs 1,00,000, Dr at source) and a matching
    YTD closing balance, so a clean run should reconcile with zero drift.
    """

    def __init__(self, branch, timeout_seconds: float = 30.0):
        self.branch = branch

    def confirm_current_company(self) -> None:
        return None

    def fetch_all_voucher_types(self, from_date, to_date):
        voucher = Voucher(
            voucher_type=VoucherType.SALES,
            voucher_date=date(2026, 1, 3),
            voucher_number="SB/1",
            branch_id=self.branch.branch_id,
            entries=[LedgerEntry(party_ledger_name="Acme", amount_as_extracted=Decimal("-100000.00"))],
        )
        return {vt: ([voucher] if vt == VoucherType.SALES else []) for vt in VoucherType}

    def fetch_ytd_sundry_debtors(self, fy_start, as_of):
        return {"Acme": Decimal("-100000.00")}


class FakeTallyClientMismatch(FakeTallyClient):
    """Same voucher stream, but the YTD closing balance disagrees with
    the roll-forward - the zero-tolerance case that must FAIL and halt.
    """

    def fetch_ytd_sundry_debtors(self, fy_start, as_of):
        return {"Acme": Decimal("-100001.00")}


def _seed_customer_master(db_path):
    store = Store(db_path)
    store.upsert_customer_master(CustomerMasterRecord("Acme", "Acme Corp", "KOL", Decimal("0.00")))
    store.close()


def test_cli_run_clean_cycle_produces_validated_report(tmp_path, monkeypatch):
    monkeypatch.setattr("ar_mis.pipeline.TallyClient", FakeTallyClient)
    monkeypatch.setattr("ar_mis.cli.TallyClient", FakeTallyClient)

    db_path = str(tmp_path / "cli.db")
    _seed_customer_master(db_path)
    report_dir = tmp_path / "reports"
    report_dir.mkdir()

    week_ending = date(2026, 1, 5)
    exit_code = cli.run(
        week_ending, db_path=db_path, report_dir=str(report_dir), branches=[KOL],
        announce=lambda *_: None, confirm=lambda *_: "",
    )
    assert exit_code == 0

    report_path = report_dir / f"AR_MIS_{week_ending.isoformat()}.xlsx"
    assert report_path.exists()
    wb = load_workbook(report_path)
    status_row = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] == "Status"][0]
    assert status_row[1] == "VALIDATED"


def test_cli_run_halts_on_reconciliation_mismatch(tmp_path, monkeypatch):
    monkeypatch.setattr("ar_mis.pipeline.TallyClient", FakeTallyClientMismatch)
    monkeypatch.setattr("ar_mis.cli.TallyClient", FakeTallyClientMismatch)

    db_path = str(tmp_path / "cli.db")
    _seed_customer_master(db_path)
    report_dir = tmp_path / "reports"
    report_dir.mkdir()

    week_ending = date(2026, 1, 5)
    exit_code = cli.run(
        week_ending, db_path=db_path, report_dir=str(report_dir), branches=[KOL],
        announce=lambda *_: None, confirm=lambda *_: "",
    )
    # Section 2.2: a reconciliation FAIL halts the run - no report at all
    # is produced for this attempt, pending manual investigation.
    assert exit_code == 1
    report_path = report_dir / f"AR_MIS_{week_ending.isoformat()}.xlsx"
    assert not report_path.exists()
