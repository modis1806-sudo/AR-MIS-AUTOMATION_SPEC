"""Section 6: static weekly snapshot report.

Client's Section 8 decision: static weekly snapshot matching the
confirmed weekly cadence, not a live/interactive dashboard. Interactive
drill-down remains a secondary tool for AR Managers, not built here.

Deliberately NOT included: day-bucket ageing (0-30/31-60/.../>180) and
the ">180-day Bad Debt Risk" KPI - two of the original defects (Section
1) were exactly a bucket sub-split not summing to its own total and a
Bad Debt Risk figure disagreeing with the Ageing Matrix for the same
date. Producing ageing buckets correctly needs bill-level due-date and
BILLTYPE (New Ref vs Against Ref) data this build does not yet extract.
Fabricating a bucket split from closing balances alone would silently
reintroduce the exact class of unvalidated-figure problem this project
exists to eliminate, so it is a named gap (see README), not a silent one.

Sheets: Summary, Party Detail, Exceptions (kept separate from Party
Detail, never blended into it - spec's explicit requirement), Movement
Trend (explicit MISSING marker for any week gap - "must not have
blank/missing weeks" is spec text, not a nice-to-have), PTP Register
(per-entry Active/Kept/Broken status).
"""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ar_mis.gate import GateStatus
from ar_mis.reconciliation import DriftFinding
from ar_mis.storage import Store

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CLEAN_FILL = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
_FAIL_FILL = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
_MISSING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")


def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autosize(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 45)


def _build_summary_sheet(wb, week_ending: date, rows_this_week, gate_status: GateStatus) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"AR MIS Weekly Report - Week Ending {week_ending.isoformat()}"])
    ws["A1"].font = Font(bold=True, size=14)

    ws.append([])
    status_text = "VALIDATED" if gate_status.clean else "UNVALIDATED - RECONCILIATION FAILED"
    ws.append(["Status", status_text])
    status_cell = ws.cell(row=ws.max_row, column=2)
    status_cell.fill = _CLEAN_FILL if gate_status.clean else _FAIL_FILL
    status_cell.font = Font(bold=True, color="FFFFFF")

    if not gate_status.clean:
        ws.append(["This report is held pending manual sign-off. Reasons:"])
        for reason in gate_status.reasons:
            ws.append(["", reason])

    ws.append([])
    total_closing = sum((Decimal(r["closing_computed"]) for r in rows_this_week), Decimal("0.00"))
    ws.append(["Consolidated Total AR (all branches, this week)", str(total_closing)])

    ws.append([])
    ws.append(["Branch", "Parties", "Reconciled", "Failed"])
    for cell in ws[ws.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    by_branch: dict[str, list] = {}
    for r in rows_this_week:
        by_branch.setdefault(r["branch_id"], []).append(r)
    for branch_id, rows in sorted(by_branch.items()):
        reconciled = sum(1 for r in rows if r["reconciled"])
        ws.append([branch_id, len(rows), reconciled, len(rows) - reconciled])

    _autosize(ws)


def _build_party_detail_sheet(wb, rows_this_week) -> None:
    ws = wb.create_sheet("Party Detail")
    _header_row(
        ws,
        [
            "Party", "Branch", "Opening", "Sales", "Credit Notes", "Debit Notes",
            "Receipts", "Journals", "Closing (Computed)", "Closing (Extracted)", "Reconciled",
        ],
    )
    by_party: dict[str, list] = {}
    for r in rows_this_week:
        by_party.setdefault(r["party_id"], []).append(r)

    for party, rows in sorted(by_party.items()):
        cross_branch_total = Decimal("0.00")
        for r in sorted(rows, key=lambda x: x["branch_id"]):
            ws.append(
                [
                    r["party_id"], r["branch_id"], r["opening"], r["sales"], r["credit_notes"],
                    r["debit_notes"], r["receipts"], r["journals"], r["closing_computed"],
                    r["closing_extracted"], "Yes" if r["reconciled"] else "No",
                ]
            )
            cross_branch_total += Decimal(r["closing_computed"])
        if len(rows) > 1:
            ws.append([party, "TOTAL (all branches)", "", "", "", "", "", "", str(cross_branch_total), "", ""])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True, italic=True)
    _autosize(ws)


def _build_exceptions_sheet(
    wb, rows_this_week, final_failed_branches, drift_findings: list[DriftFinding],
    new_parties: list[tuple[str, str]] | None = None,
) -> None:
    ws = wb.create_sheet("Exceptions")
    _header_row(ws, ["Type", "Party / Branch", "Detail"])

    for party, branch_id in new_parties or []:
        ws.append(
            [
                "New party this week",
                f"{party} / {branch_id}",
                "Auto-created with Pre-MIS Outstanding = 0 (not on the go-live seed list) - not a data error, "
                "worth a human glance.",
            ]
        )

    for r in rows_this_week:
        if not r["reconciled"]:
            ws.append(
                [
                    "Party reconciliation FAIL",
                    f"{r['party_id']} / {r['branch_id']}",
                    f"Computed {r['closing_computed']} vs Extracted {r['closing_extracted']} "
                    f"(difference {r['difference']})",
                ]
            )

    for branch in final_failed_branches:
        ws.append(["Branch extraction FAILED", branch.branch_name, branch.detail])

    for finding in drift_findings:
        week_desc = finding.attributed_week.isoformat() if finding.attributed_week else "unassigned week"
        ws.append(
            [
                "YTD cross-check drift",
                finding.party_ledger_name,
                f"{finding.voucher_type} {finding.voucher_number} dated {finding.voucher_date.isoformat()}, "
                f"amount {finding.flipped_amount}, attributed to week {week_desc}",
            ]
        )

    if ws.max_row == 1:
        ws.append(["(none)", "", ""])
    _autosize(ws)


def _expected_weeks(first: date, last: date) -> list[date]:
    weeks = []
    current = first
    while current <= last:
        weeks.append(current)
        current += timedelta(days=7)
    return weeks


def _build_movement_trend_sheet(wb, store: Store) -> None:
    ws = wb.create_sheet("Movement Trend")
    _header_row(ws, ["Party", "Branch", "Week Ending", "Closing (Computed)", "Status"])

    all_weeks = store.all_week_endings()
    if not all_weeks:
        ws.append(["(no weekly data on record yet)", "", "", "", ""])
        _autosize(ws)
        return
    latest_week = max(all_weeks)

    party_branch_pairs = {
        (rec["party_id"], rec["branch_id"]) for rec in store.all_customer_master_records()
    }
    for party_id, branch_id in sorted(party_branch_pairs):
        rows = store.weekly_snapshots_for_party(party_id, branch_id)
        if not rows:
            continue
        rows_by_week = {date.fromisoformat(r["week_ending"]): r for r in rows}
        first_week = min(rows_by_week)
        for week in _expected_weeks(first_week, latest_week):
            if week in rows_by_week:
                r = rows_by_week[week]
                ws.append([party_id, branch_id, week.isoformat(), r["closing_computed"], "OK"])
            else:
                ws.append([party_id, branch_id, week.isoformat(), "", "MISSING"])
                for cell in ws[ws.max_row]:
                    cell.fill = _MISSING_FILL
    _autosize(ws)


def _build_ptp_sheet(wb, store: Store, as_of_week: date) -> None:
    ws = wb.create_sheet("PTP Register")
    _header_row(ws, ["PTP ID", "Party", "Branch", "Promised Amount", "Promised Date", "Owner", "Status", "Notes"])
    for entry in store.all_ptp_entries():
        status = store.current_ptp_status(entry["ptp_id"], as_of_week)
        ws.append(
            [
                entry["ptp_id"], entry["party_id"], entry["branch_id"], entry["promised_amount"],
                entry["promised_date"], entry["owner"], status.value if status else "(no status logged)",
                entry["notes"],
            ]
        )
    _autosize(ws)


def generate_report(
    store: Store,
    week_ending: date,
    gate_status: GateStatus,
    final_failed_branches,
    drift_findings: list[DriftFinding],
    output_path: str,
    new_parties: list[tuple[str, str]] | None = None,
) -> None:
    """Always writes a report file - a non-clean gate does not mean "no
    report", it means the report is stamped UNVALIDATED and held for
    manual sign-off rather than auto-sent (see ar_mis.gate.may_auto_send).
    """
    rows_this_week = store.weekly_snapshots_for_week(week_ending)

    wb = Workbook()
    _build_summary_sheet(wb, week_ending, rows_this_week, gate_status)
    _build_party_detail_sheet(wb, rows_this_week)
    _build_exceptions_sheet(wb, rows_this_week, final_failed_branches, drift_findings, new_parties)
    _build_movement_trend_sheet(wb, store)
    _build_ptp_sheet(wb, store, week_ending)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2" if ws.title != "Summary" else None
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.replace(".", "", 1).lstrip("-").isdigit():
                    cell.alignment = Alignment(horizontal="right")

    wb.save(output_path)
